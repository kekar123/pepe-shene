from __future__ import annotations

import json
import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

import pandas as pd


def _normalize(text: str) -> str:
    if text is None:
        return ""
    value = str(text).strip().lower()
    # Keep only alphanumeric symbols (works for latin/cyrillic and avoids mojibake regex issues).
    return "".join(ch for ch in value if ch.isalnum())


def _find_column(columns: Iterable, candidates: List[str]) -> Optional[str]:
    normalized = {}
    for col in columns:
        norm = _normalize(col)
        if norm:
            normalized[norm] = col
    for cand in candidates:
        cand_norm = _normalize(cand)
        if not cand_norm:
            continue
        if cand_norm in normalized:
            return normalized[cand_norm]
    # fallback: partial match
    for cand in candidates:
        cand_norm = _normalize(cand)
        if not cand_norm:
            continue
        for norm_col, original in normalized.items():
            if cand_norm and cand_norm in norm_col:
                return original
    return None


def _column_by_index(columns: Iterable, index: int) -> Optional[str]:
    cols = list(columns)
    if 0 <= index < len(cols):
        return cols[index]
    return None


def _is_reappro_reason(value) -> bool:
    """
    Reappro includes pallet drop/pallet preparation operations.
    """
    norm = _normalize(value)
    if not norm:
        return False
    has_pallet = "пал" in norm or "pallet" in norm
    is_drop = "спуск" in norm or "drop" in norm
    is_prepare = "подготов" in norm or "prepare" in norm
    return has_pallet and (is_drop or is_prepare)


def _clean_article(value) -> Optional[str]:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    text = str(value).strip()
    return text if text else None


def _read_excel(path: Path, sheet_name=0, header=0) -> pd.DataFrame:
    return pd.read_excel(path, sheet_name=sheet_name, header=header)


def _parse_location(location: str) -> Tuple[Optional[str], Optional[str]]:
    if not location:
        return None, None
    parts = str(location).strip().split("-")
    parts = [p for p in parts if p != ""]
    if len(parts) >= 3:
        aisle = parts[1]
        place = parts[2]
        return aisle, place
    # fallback: try to extract groups of digits
    digits = re.findall(r"\d+", str(location))
    if len(digits) >= 2:
        return digits[0], digits[1]
    return None, None


def _is_zero_like_location(value) -> bool:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return False
    if isinstance(value, (int, float)):
        return float(value) == 0.0
    text = str(value).strip()
    if not text:
        return False
    if set(text) == {"0"}:
        return True
    digits = re.findall(r"\d+", text)
    if digits and all(int(d) == 0 for d in digits):
        return True
    return False


def _load_master_data(path: Path) -> pd.DataFrame:
    df = _read_excel(path, sheet_name=0, header=1)
    code_col = _find_column(df.columns, ["Код артикула", "Article Code"])
    name_col = _find_column(df.columns, ["Имя артикула", "Article Name"])
    pcb_col = _find_column(df.columns, ["Шт в коробе", "Pieces Per Case"])
    stockage_col = _find_column(df.columns, ["Класс стокажа", "Storage Class"])
    weight_col = _find_column(df.columns, ["Weight", "Вес", "Вес\\Weight", "Weight\\Вес"])
    # Fallbacks by expected layout from master file (C,D,M,J,T columns).
    code_col = code_col or _column_by_index(df.columns, 2)
    name_col = name_col or _column_by_index(df.columns, 3)
    pcb_col = pcb_col or _column_by_index(df.columns, 12)
    weight_col = weight_col or _column_by_index(df.columns, 9)
    stockage_col = stockage_col or _column_by_index(df.columns, 19)
    cols = {}
    if code_col:
        cols["article"] = df[code_col].apply(_clean_article)
    if name_col:
        cols["name"] = df[name_col]
    if pcb_col:
        cols["pcb"] = df[pcb_col]
    if stockage_col:
        cols["stockage"] = df[stockage_col]
    if weight_col:
        cols["weight"] = df[weight_col]
    result = pd.DataFrame(cols)
    if "article" not in result.columns:
        return pd.DataFrame(columns=["article", "name", "pcb", "stockage", "weight"])
    result = result.dropna(subset=["article"])
    # РЈР±РёСЂР°РµРј СЃС‚СЂРѕРєСѓ-С€Р°РїРєСѓ РЅР° Р°РЅРіР»РёР№СЃРєРѕРј, РµСЃР»Рё РїРѕРїР°Р»Р° РІ РґР°РЅРЅС‹Рµ
    result = result[result["article"].astype(str).str.strip().str.lower().ne("article code")]
    return result


def _load_stock_report(path: Path) -> pd.DataFrame:
    df = _read_excel(path, sheet_name=0, header=0)
    article_col = _find_column(df.columns, ["Article", "Артикул"])
    location_col = _find_column(df.columns, ["Location", "Локация", "Место"])

    if not article_col or not location_col:
        return pd.DataFrame(columns=["article", "location", "aisle", "place"])

    data = df[[article_col, location_col]].copy()
    data.columns = ["article", "location"]
    data["article"] = data["article"].apply(_clean_article)
    data = data.dropna(subset=["article"])

    # choose most frequent location per article
    location_mode = (
        data.groupby("article")["location"]
        .agg(lambda x: x.dropna().astype(str).mode().iloc[0] if not x.dropna().empty else None)
        .reset_index()
    )
    location_mode["aisle"] = location_mode["location"].apply(lambda x: _parse_location(x)[0])
    location_mode["place"] = location_mode["location"].apply(lambda x: _parse_location(x)[1])
    return location_mode


def _load_order_picked(path: Path) -> pd.DataFrame:
    try:
        df = _read_excel(path, sheet_name="За период", header=4)
    except Exception:
        df = _read_excel(path, sheet_name=0, header=4)
    # remove service row with codes
    if "Заказ" in df.columns:
        df = df[df["Заказ"] != "REFLIV"]
    article_col = _find_column(df.columns, ["Артикул", "АРТИКУЛ"])
    picked_boxes_col = _find_column(df.columns, ["Собрано кор.", "Собрано кор", "Собрано"])
    aisle_col = _find_column(df.columns, ["Аллея"])
    place_col = _find_column(df.columns, ["Место"])
    # Fallbacks by expected layout from order-picked file (D,H,M,N columns).
    article_col = article_col or _column_by_index(df.columns, 3)
    picked_boxes_col = picked_boxes_col or _column_by_index(df.columns, 7)
    aisle_col = aisle_col or _column_by_index(df.columns, 12)
    place_col = place_col or _column_by_index(df.columns, 13)

    data = pd.DataFrame()
    if article_col:
        data["article"] = df[article_col].apply(_clean_article)
    if picked_boxes_col:
        data["picked_boxes"] = pd.to_numeric(df[picked_boxes_col], errors="coerce").fillna(0)
    if aisle_col:
        data["aisle"] = df[aisle_col]
    if place_col:
        data["place"] = df[place_col]

    if "article" not in data.columns:
        return pd.DataFrame(columns=["article", "max_pick", "lines_count", "aisle", "place"])
    data = data.dropna(subset=["article"])

    if "aisle" not in data.columns:
        data["aisle"] = None
    if "place" not in data.columns:
        data["place"] = None

    grouped = data.groupby("article").agg(
        max_pick=("picked_boxes", "sum"),
        lines_count=("article", "count"),
        aisle=("aisle", lambda x: x.dropna().astype(str).mode().iloc[0] if not x.dropna().empty else None),
        place=("place", lambda x: x.dropna().astype(str).mode().iloc[0] if not x.dropna().empty else None),
    ).reset_index()
    return grouped


def _load_stock_movements(path: Path) -> pd.DataFrame:
    df = _read_excel(path, sheet_name=0, header=0)
    reason_col = _find_column(df.columns, ["Причина"])
    article_col = _find_column(df.columns, ["Артикул"])
    pallets_col = _find_column(df.columns, ["Номер паллета", "SSCC"])
    boxes_col = _find_column(df.columns, ["Коробов", "Коробки"])
    units_col = _find_column(df.columns, ["Штук", "ШТУК"])
    # Fallbacks by expected layout from stock movements file.
    reason_col = reason_col or _column_by_index(df.columns, 1)
    article_col = article_col or _column_by_index(df.columns, 12)
    pallets_col = pallets_col or _column_by_index(df.columns, 8) or _column_by_index(df.columns, 7)
    units_col = units_col or _column_by_index(df.columns, 15)
    boxes_col = boxes_col or _column_by_index(df.columns, 16)

    if not reason_col or not article_col:
        return pd.DataFrame(columns=["article", "reappro_pallets", "reappro_boxes", "reappro_units"])

    df = df.copy()
    df["article"] = df[article_col].apply(_clean_article)
    df["reason"] = df[reason_col].astype(str).str.lower()
    df = df.dropna(subset=["article"])

    filtered = df[df["reason"].apply(_is_reappro_reason)].copy()
    if filtered.empty:
        return pd.DataFrame(columns=["article", "reappro_pallets", "reappro_boxes", "reappro_units"])

    if pallets_col and pallets_col in filtered.columns:
        pallet_series = filtered[pallets_col]
    else:
        pallet_series = pd.Series([None] * len(filtered), index=filtered.index)

    filtered["reappro_pallets"] = pallet_series
    filtered["reappro_boxes"] = pd.to_numeric(filtered[boxes_col], errors="coerce").fillna(0) if boxes_col else 0
    filtered["reappro_units"] = pd.to_numeric(filtered[units_col], errors="coerce").fillna(0) if units_col else 0

    def pallet_count(series: pd.Series) -> int:
        series = series.dropna()
        if series.empty:
            return 0
        return series.nunique()

    grouped = filtered.groupby("article").agg(
        reappro_pallets=("reappro_pallets", pallet_count),
        reappro_boxes=("reappro_boxes", "sum"),
        reappro_units=("reappro_units", "sum"),
    ).reset_index()

    return grouped


def _load_picking_lines(path: Path) -> pd.DataFrame:
    df = _read_excel(path, sheet_name=0, header=0)
    article_col = _find_column(df.columns, ["АРТИКУЛ", "Артикул"])
    order_col = _find_column(df.columns, ["ЗАКАЗ", "Заказ"])
    date_col = _find_column(df.columns, ["ПЛАНОВАЯ_ДАТА_ОТГРУЗКИ", "Дата отгрузки"])
    support_type_col = _find_column(df.columns, ["ТИП_СУППОРТА", "Тип суппорта"])
    picked_boxes_col = _find_column(df.columns, ["СОБРАНО_КОР", "Собрано кор"])
    picked_units_col = _find_column(df.columns, ["ШТ_ДЛЯ_СБОРКИ", "Собрано шт", "ШТУК_ЗАКАЗАНО"])
    # Fallbacks by expected layout from picking lines file.
    article_col = article_col or _column_by_index(df.columns, 12)
    order_col = order_col or _column_by_index(df.columns, 2)
    date_col = date_col or _column_by_index(df.columns, 3)
    support_type_col = support_type_col or _column_by_index(df.columns, 6)
    picked_boxes_col = picked_boxes_col or _column_by_index(df.columns, 19)
    picked_units_col = picked_units_col or _column_by_index(df.columns, 21) or _column_by_index(df.columns, 20)

    if not article_col:
        return pd.DataFrame(columns=["article", "orders_count", "first_out_date", "last_out_date", "picked_boxes", "picked_units", "pick_days_count"])

    df = df.copy()
    df["article"] = df[article_col].apply(_clean_article)
    df = df.dropna(subset=["article"])

    if support_type_col and support_type_col in df.columns:
        df["support_type"] = df[support_type_col].astype(str).str.lower()
        df_pick = df[df["support_type"].str.contains("пик", na=False)].copy()
    else:
        df_pick = df.copy()

    df_pick["picked_boxes"] = pd.to_numeric(df_pick[picked_boxes_col], errors="coerce").fillna(0) if picked_boxes_col else 0
    df_pick["picked_units"] = pd.to_numeric(df_pick[picked_units_col], errors="coerce").fillna(0) if picked_units_col else 0

    df["orders_count"] = df[order_col] if order_col else None
    df["last_out_date"] = pd.to_datetime(df[date_col], errors="coerce") if date_col else pd.NaT

    orders = (
        df.groupby("article")["orders_count"]
        .agg(lambda x: x.dropna().nunique())
        .reset_index()
        .rename(columns={"orders_count": "orders_count"})
    )

    date_stats = (
        df.groupby("article")["last_out_date"]
        .agg(first_out_date="min", last_out_date="max")
        .reset_index()
    )

    # Р”Р°С‚С‹ РїРѕРґР±РѕСЂР° РІ РєРѕСЂРѕР±Р°С… (Р±РµСЂРµРј С‚РѕР»СЊРєРѕ СЃС‚СЂРѕРєРё, РіРґРµ СЃРѕР±СЂР°РЅС‹ РєРѕСЂРѕР±Р°)
    if date_col:
        df_pick_dates = df_pick.copy()
        df_pick_dates["pick_date"] = pd.to_datetime(df_pick_dates[date_col], errors="coerce")
        df_pick_dates["pick_boxes"] = pd.to_numeric(df_pick_dates[picked_boxes_col], errors="coerce").fillna(0) if picked_boxes_col else 0
        df_pick_dates = df_pick_dates[df_pick_dates["pick_boxes"] > 0]
        pick_days = (
            df_pick_dates.groupby("article")["pick_date"]
            .agg(
                first_pick_date="min",
                last_pick_date="max",
                pick_days_count=lambda x: x.dropna().dt.date.nunique()
            )
            .reset_index()
        )
    else:
        pick_days = pd.DataFrame(columns=["article", "first_pick_date", "last_pick_date", "pick_days_count"])

    picks = (
        df_pick.groupby("article")
        .agg(
            picked_boxes=("picked_boxes", "sum"),
            picked_units=("picked_units", "sum"),
        )
        .reset_index()
    )

    combined = (
        orders.merge(date_stats, on="article", how="outer")
        .merge(picks, on="article", how="outer")
        .merge(pick_days, on="article", how="outer")
    )
    return combined


def _load_abc_analysis(path: Path) -> pd.DataFrame:
    df = _read_excel(path, sheet_name=0, header=0)
    # РЈР±РёСЂР°РµРј СЃС‚СЂРѕРєСѓ-С€Р°РїРєСѓ, РµСЃР»Рё РѕРЅР° РїСЂРѕРґСѓР±Р»РёСЂРѕРІР°РЅР° РєР°Рє РїРµСЂРІР°СЏ СЃС‚СЂРѕРєР° РґР°РЅРЅС‹С…
    if not df.empty:
        first_row = df.iloc[0].tolist()
        if any("артикул" in str(cell).strip().lower() for cell in first_row):
            df = df.iloc[1:].reset_index(drop=True)
    article_col = _find_column(df.columns, ["Артикул", "АРТИКУЛ"])
    total_boxes_col = _find_column(df.columns, ["Коробов всего", "Коробов всего "])
    palletization_col = _find_column(df.columns, ["Паллетизация"])
    # Fallbacks by expected layout from ABC file (A,I,O columns).
    article_col = article_col or _column_by_index(df.columns, 0)
    total_boxes_col = total_boxes_col or _column_by_index(df.columns, 8)
    palletization_col = palletization_col or _column_by_index(df.columns, 14)

    data = pd.DataFrame()
    if article_col:
        data["article"] = df[article_col].apply(_clean_article)
    if total_boxes_col:
        data["total_boxes"] = pd.to_numeric(df[total_boxes_col], errors="coerce").fillna(0)
    if palletization_col:
        data["palletization"] = df[palletization_col]

    if "article" not in data.columns:
        return pd.DataFrame(columns=["article", "total_boxes", "palletization"])
    data = data.dropna(subset=["article"])
    grouped = data.groupby("article").agg(
        total_boxes=("total_boxes", "sum"),
        palletization=("palletization", lambda x: x.dropna().iloc[0] if not x.dropna().empty else None),
    ).reset_index()
    return grouped


def _compute_abc_classes(df: pd.DataFrame, value_col: str) -> pd.DataFrame:
    data = df.copy()
    data[value_col] = pd.to_numeric(data[value_col], errors="coerce").fillna(0)
    data = data.sort_values(value_col, ascending=False)
    total = data[value_col].sum()
    if total <= 0:
        data["abc_class"] = "C"
        return data[["article", "abc_class"]]

    data["share"] = data[value_col] / total * 100
    data["cumulative"] = data["share"].cumsum()

    def classify(cum):
        if cum <= 80:
            return "A"
        if cum <= 95:
            return "B"
        return "C"

    data["abc_class"] = data["cumulative"].apply(classify)
    return data[["article", "abc_class"]]


def _compute_abc_classes_by_frequency(picking_lines: pd.DataFrame) -> pd.DataFrame:
    data = picking_lines.copy()
    if data.empty or "article" not in data.columns:
        return pd.DataFrame(columns=["article", "abc_class"])

    data["pick_days_count"] = pd.to_numeric(data.get("pick_days_count"), errors="coerce").fillna(0)
    data["first_pick_date"] = pd.to_datetime(data.get("first_pick_date"), errors="coerce")
    data["last_pick_date"] = pd.to_datetime(data.get("last_pick_date"), errors="coerce")

    global_min = data["first_pick_date"].min()
    global_max = data["last_pick_date"].max()
    if pd.isna(global_min) or pd.isna(global_max):
        global_period_days = 0
    else:
        global_period_days = max((global_max.date() - global_min.date()).days, 0)

    def avg_interval_days(row) -> float:
        pick_days = float(row.get("pick_days_count") or 0)
        first = row.get("first_pick_date")
        last = row.get("last_pick_date")

        if pick_days <= 1:
            # Р•СЃР»Рё РїРѕРґР±РѕСЂ РІ РєРѕСЂРѕР±Р°С… Р±С‹Р» РѕРґРёРЅ СЂР°Р·, Р±РµСЂРµРј РѕР±С‰РёР№ РїРµСЂРёРѕРґ РІС‹РіСЂСѓР·РєРё.
            return float(global_period_days or 9999)

        if pd.isna(first) or pd.isna(last):
            return float(global_period_days or 9999)

        period_days = max((last.date() - first.date()).days, 0)
        intervals = max(pick_days - 1, 1)
        return period_days / intervals if intervals else float(global_period_days or 9999)

    data["avg_interval_days"] = data.apply(avg_interval_days, axis=1)

    def classify(days: float) -> str:
        if days <= 14:
            return "A"
        if days <= 30:
            return "B"
        return "C"

    data["abc_class"] = data["avg_interval_days"].apply(classify)
    return data[["article", "abc_class"]]


def _compute_popularity(df: pd.DataFrame, value_col: str, percent_col: str, popularity_col: str) -> pd.DataFrame:
    data = df[["article", value_col]].copy()
    data[value_col] = pd.to_numeric(data[value_col], errors="coerce").fillna(0)
    data = data.sort_values(value_col, ascending=False)
    total = data[value_col].sum()
    if total <= 0:
        data[percent_col] = 0.0
        data[popularity_col] = 0.0
        return data[["article", percent_col, popularity_col]]

    data[percent_col] = data[value_col] / total * 100
    data[popularity_col] = data[percent_col].cumsum()
    return data[["article", percent_col, popularity_col]]



def _parse_weight_kg(value) -> Optional[float]:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if isinstance(value, (int, float)):
        return float(value)

    text = str(value).strip().replace(",", ".")
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    if not match:
        return None
    try:
        return float(match.group(0))
    except ValueError:
        return None


def _weight_abc_class(value) -> Optional[str]:
    weight = _parse_weight_kg(value)
    if weight is None:
        return None
    if 0.1 <= weight <= 1.5:
        return "C"
    if 1.6 <= weight <= 5.0:
        return "B"
    if 5.1 <= weight <= 20.0:
        return "A"
    return None


def _dead_stock_status(last_out_date: Optional[pd.Timestamp]) -> str:
    if last_out_date is None or pd.isna(last_out_date):
        return "нет данных"
    days = (datetime.now().date() - last_out_date.date()).days
    if days >= 90:
        return "красный"
    if days >= 60:
        return "желтый"
    return "зеленый"


def _dead_stock_sort_rank(last_out_date: Optional[pd.Timestamp]) -> int:
    if last_out_date is None or pd.isna(last_out_date):
        return 3
    days = (datetime.now().date() - last_out_date.date()).days
    if days >= 90:
        return 2
    if days >= 60:
        return 1
    return 0


def _safe_str(value: Any) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    return str(value).strip()


def _safe_num(value: Any) -> float:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return 0.0
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def _build_ai_summary(master: pd.DataFrame, output: pd.DataFrame) -> Dict[str, Any]:
    """Формирует структуру для скрытого JSON отчёта (для последующего использования нейросетью)."""
    count_master = len(master) if master is not None and not master.empty else 0
    count_picking = len(output) if output is not None and not output.empty else 0

    abc_freq: Dict[str, int] = {}
    abc_weight: Dict[str, int] = {}
    if output is not None and not output.empty:
        if "ABC класс" in output.columns:
            for k, v in output["ABC класс"].value_counts().items():
                key = _safe_str(k)
                if key:
                    abc_freq[key] = int(v)
        if "ABC класс (вес)" in output.columns:
            for k, v in output["ABC класс (вес)"].value_counts().items():
                key = _safe_str(k)
                if key:
                    abc_weight[key] = int(v)

    def top10(sort_col: str, descending: bool = True) -> List[Dict[str, Any]]:
        if output is None or output.empty or sort_col not in output.columns:
            return []
        df = output.copy()
        df["_val"] = df[sort_col].apply(_safe_num)
        df = df[df["_val"] > 0].sort_values("_val", ascending=not descending).head(10)
        return [
            {
                "артикул": _safe_str(row.get("АРТИКУЛ")),
                "название": _safe_str(row.get("НАЗВАНИЕ")),
                "значение": _safe_num(row.get(sort_col)),
            }
            for _, row in df.iterrows()
        ]

    top10_orders = top10("Кол-во заказов")
    top10_picked_units = top10("Выход в штуках")
    top10_lines = top10("Кол-во линий")
    top10_reappro = top10("Кол-во коробов реапро")

    # Топ-10 худших: только строки с данными (есть хотя бы заказы, выход или линии), сортировка по возрастанию
    top10_worst: List[Dict[str, Any]] = []
    if output is not None and not output.empty:
        orders_col = output.get("Кол-во заказов", pd.Series(dtype=float))
        units_col = output.get("Выход в штуках", pd.Series(dtype=float))
        lines_col = output.get("Кол-во линий", pd.Series(dtype=float))
        if orders_col is None:
            orders_col = pd.Series(0.0, index=output.index)
        if units_col is None:
            units_col = pd.Series(0.0, index=output.index)
        if lines_col is None:
            lines_col = pd.Series(0.0, index=output.index)
        has_data = (orders_col.fillna(0) > 0) | (units_col.fillna(0) > 0) | (lines_col.fillna(0) > 0)
        with_data = output.loc[has_data].copy()
        with_data["_score"] = (
            with_data["Кол-во заказов"].fillna(0) + with_data["Выход в штуках"].fillna(0) * 0.001 + with_data["Кол-во линий"].fillna(0)
        )
        with_data = with_data.sort_values("_score", ascending=True).head(10)
        for _, row in with_data.iterrows():
            top10_worst.append({
                "артикул": _safe_str(row.get("АРТИКУЛ")),
                "название": _safe_str(row.get("НАЗВАНИЕ")),
                "кол_во_заказов": _safe_num(row.get("Кол-во заказов")),
                "выход_в_штуках": _safe_num(row.get("Выход в штуках")),
                "кол_во_линий": _safe_num(row.get("Кол-во линий")),
            })

    # Топ-10 залежавшихся: по дате последнего выхода (старые первые), с количеством дней
    top10_stagnant: List[Dict[str, Any]] = []
    if output is not None and not output.empty and "Дата последнего выхода" in output.columns:
        today = datetime.now().date()
        rows_with_date = []
        for _, row in output.iterrows():
            raw = row.get("Дата последнего выхода")
            if raw is None or (isinstance(raw, float) and pd.isna(raw)):
                continue
            if isinstance(raw, pd.Timestamp):
                dt = raw.date()
            else:
                try:
                    dt = datetime.strptime(str(raw).strip()[:10], "%Y-%m-%d").date()
                except (ValueError, TypeError):
                    continue
            days_lying = (today - dt).days
            rows_with_date.append((days_lying, dt, row))
        rows_with_date.sort(key=lambda x: -x[0])
        for days_lying, dt, row in rows_with_date[:10]:
            top10_stagnant.append({
                "артикул": _safe_str(row.get("АРТИКУЛ")),
                "название": _safe_str(row.get("НАЗВАНИЕ")),
                "дата_последнего_выхода": dt.strftime("%Y-%m-%d"),
                "дней_лежит": days_lying,
            })

    return {
        "количество_артикулов_мастер": count_master,
        "количество_товаров_в_пикинге": count_picking,
        "abc_по_частоте": abc_freq,
        "abc_по_весу": abc_weight,
        "топ10_по_заказам": top10_orders,
        "топ10_по_выходу_в_штуках": top10_picked_units,
        "топ10_по_количеству_линий": top10_lines,
        "топ10_по_коробам_реапро": top10_reappro,
        "топ10_худших": top10_worst,
        "топ10_залежавшихся": top10_stagnant,
    }


@dataclass
class CombinedReportResult:
    rows: List[Dict]
    file_path: Path


def generate_combined_report(
    master_file: Path,
    stock_report_file: Path,
    order_picked_file: Path,
    stock_movements_file: Path,
    picking_lines_file: Path,
    abc_analysis_file: Path,
    output_dir: Path,
) -> CombinedReportResult:
    master = _load_master_data(master_file)
    stock_report = _load_stock_report(stock_report_file)
    order_picked = _load_order_picked(order_picked_file)
    stock_movements = _load_stock_movements(stock_movements_file)
    picking_lines = _load_picking_lines(picking_lines_file)
    abc_data = _load_abc_analysis(abc_analysis_file)

    all_articles = pd.Series(
        pd.concat([
            master["article"] if "article" in master else pd.Series(dtype=str),
            stock_report["article"] if "article" in stock_report else pd.Series(dtype=str),
            order_picked["article"] if "article" in order_picked else pd.Series(dtype=str),
            stock_movements["article"] if "article" in stock_movements else pd.Series(dtype=str),
            picking_lines["article"] if "article" in picking_lines else pd.Series(dtype=str),
            abc_data["article"] if "article" in abc_data else pd.Series(dtype=str),
        ], ignore_index=True)
    ).dropna().drop_duplicates()

    base = pd.DataFrame({"article": all_articles})

    report = base.merge(master, on="article", how="left")
    report = report.merge(stock_report[["article", "aisle", "place"]], on="article", how="left")
    report = report.merge(order_picked[["article", "max_pick", "lines_count", "aisle", "place"]], on="article", how="left", suffixes=("", "_order"))
    report = report.merge(stock_movements, on="article", how="left")
    report = report.merge(picking_lines, on="article", how="left")
    report = report.merge(abc_data, on="article", how="left")

    # fill aisle/place from order-picked if missing
    report["aisle"] = report["aisle"].fillna(report.get("aisle_order"))
    report["place"] = report["place"].fillna(report.get("place_order"))

    abc_classes = _compute_abc_classes_by_frequency(picking_lines)
    report = report.merge(abc_classes, on="article", how="left")

    report_filled = report.fillna(0).infer_objects(copy=False)
    reappro_pop = _compute_popularity(report_filled, "reappro_boxes", "percent_reappro", "popularity_reappro")
    lines_pop = _compute_popularity(report_filled, "lines_count", "percent_lines", "popularity_lines")

    report = report.merge(reappro_pop, on="article", how="left")
    report = report.merge(lines_pop, on="article", how="left")
    report["weight_abc_class"] = report.get("weight").apply(_weight_abc_class)
    report["dead_stock_status"] = report["last_out_date"].apply(_dead_stock_status)
    report["dead_stock_rank"] = report["last_out_date"].apply(_dead_stock_sort_rank)
    report = report.sort_values(by=["dead_stock_rank", "article"], ascending=[True, True]).reset_index(drop=True)

    # Remove rows where picking place is zero-like.
    place_series = report.get("place")
    if place_series is not None:
        report = report.loc[~place_series.apply(_is_zero_like_location)].reset_index(drop=True)

    # output columns
    output = pd.DataFrame({
        "АРТИКУЛ": report["article"],
        "НАЗВАНИЕ": report.get("name"),
        "Аллея пикинг": report.get("aisle"),
        "Место пикинг": report.get("place"),
        "MAX PICK": report.get("max_pick", 0).fillna(0),
        "Кол-во палет реапро": report.get("reappro_pallets", 0).fillna(0),
        "Кол-во коробов реапро": report.get("reappro_boxes", 0).fillna(0),
        "Кол-во штук Реапро": report.get("reappro_units", 0).fillna(0),
        "Кол-во заказов": report.get("orders_count", 0).fillna(0),
        "Дата последнего выхода": report.get("last_out_date"),
        "Выход в коробах из пикинга": report.get("picked_boxes", 0).fillna(0),
        "Выход в штуках": report.get("picked_units", 0).fillna(0),
        "Палетизация": report.get("palletization"),
        "PCB": report.get("pcb"),
        "Вес": report.get("weight"),
        "ABC класс": report.get("abc_class"),
        "ABC класс (вес)": report.get("weight_abc_class"),
        "Стокаж": report.get("stockage"),
        "Кол-во линий": report.get("lines_count", 0).fillna(0),
        "Процент (реаппро)": report.get("percent_reappro", 0).fillna(0),
        "Популярность (реаппро)": report.get("popularity_reappro", 0).fillna(0),
        "Процент (линии)": report.get("percent_lines", 0).fillna(0),
        "Популярность (линии)": report.get("popularity_lines", 0).fillna(0),
        "Мертвый сток": report.get("dead_stock_status"),
    })

    if "Дата последнего выхода" in output.columns:
        output["Дата последнего выхода"] = output["Дата последнего выхода"].apply(
            lambda value: value.strftime("%Y-%m-%d") if isinstance(value, pd.Timestamp) else ("" if pd.isna(value) else str(value))
        )
    for percent_col in ["Процент (реаппро)", "Популярность (реаппро)", "Процент (линии)", "Популярность (линии)"]:
        if percent_col in output.columns:
            output[percent_col] = pd.to_numeric(output[percent_col], errors="coerce").fillna(0).round(2)

    output_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    file_path = output_dir / f"combined_report_{timestamp}.xlsx"
    json_path = output_dir / f"combined_report_{timestamp}.json"
    try:
        ai_summary = _build_ai_summary(master, output)
        with open(json_path, "w", encoding="utf-8") as f:
            json.dump(ai_summary, f, ensure_ascii=False, indent=2, default=str)
    except Exception:
        pass

    with pd.ExcelWriter(file_path) as writer:
        output.to_excel(writer, index=False, sheet_name="Отчет")
        worksheet = writer.sheets["Отчет"]

        # РЁРёСЂРёРЅС‹ РєРѕР»РѕРЅРѕРє
        worksheet.column_dimensions["A"].width = 18   # РђСЂС‚РёРєСѓР»
        worksheet.column_dimensions["B"].width = 45   # РќР°Р·РІР°РЅРёРµ
        worksheet.column_dimensions["C"].width = 16
        worksheet.column_dimensions["D"].width = 16
        worksheet.column_dimensions["E"].width = 12
        worksheet.column_dimensions["F"].width = 16
        worksheet.column_dimensions["G"].width = 16
        worksheet.column_dimensions["H"].width = 16
        worksheet.column_dimensions["I"].width = 18
        worksheet.column_dimensions["J"].width = 18
        worksheet.column_dimensions["K"].width = 18
        worksheet.column_dimensions["L"].width = 18
        worksheet.column_dimensions["M"].width = 14
        worksheet.column_dimensions["N"].width = 14
        worksheet.column_dimensions["O"].width = 12
        worksheet.column_dimensions["P"].width = 14
        worksheet.column_dimensions["Q"].width = 14
        worksheet.column_dimensions["R"].width = 14
        worksheet.column_dimensions["S"].width = 18
        worksheet.column_dimensions["T"].width = 18
        worksheet.column_dimensions["U"].width = 18
        worksheet.column_dimensions["V"].width = 18
        worksheet.column_dimensions["W"].width = 16

        # РџРµСЂРµРЅРѕСЃ С‚РµРєСЃС‚Р° Рё РІС‹СЃРѕС‚Р° СЃС‚СЂРѕРє
        from openpyxl.styles import Alignment, Font, PatternFill
        header_fill = PatternFill(fill_type="solid", start_color="00008B", end_color="00008B")
        header_font = Font(color="FFFFFF", bold=True)
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
        dead_stock_fills = {
            "red": PatternFill(fill_type="solid", start_color="FFFF0000", end_color="FFFF0000"),
            "yellow": PatternFill(fill_type="solid", start_color="FFFFFF00", end_color="FFFFFF00"),
            "green": PatternFill(fill_type="solid", start_color="FF00B050", end_color="FF00B050"),
        }

        # РџРѕРґСЃРІРµС‚РєР° Р°СЂС‚РёРєСѓР»Р° (РєРѕР»РѕРЅРєР° A) РїРѕ СЃС‚Р°С‚СѓСЃСѓ "РњРµСЂС‚РІС‹Р№ СЃС‚РѕРє" (РїРѕСЃР»РµРґРЅСЏСЏ РєРѕР»РѕРЅРєР°)
        for row_idx, last_out_date in enumerate(report.get("last_out_date", pd.Series(dtype="datetime64[ns]")), start=2):
            if last_out_date is None or pd.isna(last_out_date):
                continue
            days = (datetime.now().date() - last_out_date.date()).days
            color_key = "red" if days >= 90 else ("yellow" if days >= 60 else "green")
            worksheet.cell(row=row_idx, column=1).fill = dead_stock_fills[color_key]

        wrap = Alignment(wrap_text=True, horizontal="center", vertical="center")
        for row in worksheet.iter_rows():
            worksheet.row_dimensions[row[0].row].height = 30
            for cell in row:
                cell.alignment = wrap

    safe_output = output.where(pd.notna(output), None)
    rows = safe_output.to_dict(orient="records")
    return CombinedReportResult(rows=rows, file_path=file_path)

